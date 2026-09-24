from __future__ import annotations

from copy import deepcopy
from datetime import date, timedelta
from typing import Any

import pandas as pd

from modules.config import EMPTY_LABEL, EMPTY_TOKEN
from modules.data_loader import (
    load_taiwan_holidays,
    profile_column_for_weight,
    tokens_from_sequence_row,
    test_info_by_id,
)

BLOCKED_FILL = "#ffcdd2"  # light red for weekend / holiday columns
EVENT_MILESTONE_FILL = "#fff9c4"  # light yellow for System ETA / Critical Feedback


def format_date_display(d: date | str) -> str:
    """Display dates as 2026/09/28."""
    if isinstance(d, date):
        return d.strftime("%Y/%m/%d")
    text = str(d).strip()
    if len(text) >= 10 and text[4] == "-" and text[7] == "-":
        return f"{text[0:4]}/{text[5:7]}/{text[8:10]}"
    return text.replace("-", "/")


def daterange(start: date, end: date) -> list[date]:
    days = (end - start).days
    if days < 0:
        return []
    return [start + timedelta(days=i) for i in range(days + 1)]


def _holiday_map_for_span(start: date, end: date) -> dict[str, str]:
    """
    Load holiday years that actually cover [start, end], plus a modest day pad
    for business-day look-ahead — not a full ±1 calendar year (that pulled 2028
    when planning Dec 2026 → early 2027, and 2028 is often unpublished).
    """
    if end < start:
        start, end = end, start
    pad_start = start - timedelta(days=14)
    pad_end = end + timedelta(days=90)
    years = tuple(range(pad_start.year, pad_end.year + 1))
    return load_taiwan_holidays(years)


def is_non_working(d: date, holidays: dict[str, str] | None = None) -> bool:
    if d.weekday() >= 5:
        return True
    if holidays is None:
        holidays = load_taiwan_holidays((d.year,))
    return d.isoformat() in holidays


def add_business_days(d: date, n: int, holidays: dict[str, str] | None = None) -> date:
    """Move forward/backward by n Taiwan business days (skip weekends + holidays)."""
    if n == 0:
        return d
    if holidays is None:
        # rough year window
        approx_end = d + timedelta(days=abs(n) * 3 + 14)
        holidays = _holiday_map_for_span(min(d, approx_end), max(d, approx_end))
    step = 1 if n > 0 else -1
    remaining = abs(n)
    cur = d
    while remaining > 0:
        cur += timedelta(days=step)
        if not is_non_working(cur, holidays):
            remaining -= 1
    return cur


def compute_timeline_end(
    system_eta: date,
    critical_feedback: date,
    plan_end: date | None = None,
    holidays: dict[str, str] | None = None,
) -> date:
    """
    End = max(
      System ETA,
      5 business days after Critical Feedback,
      final test-plan occupied day (if later),
    )
    """
    if holidays is None:
        probe = max(system_eta, critical_feedback) + timedelta(days=40)
        holidays = _holiday_map_for_span(min(system_eta, critical_feedback), probe)
    cf_plus_5 = add_business_days(critical_feedback, 5, holidays)
    end = max(system_eta, cf_plus_5)
    if plan_end is not None:
        end = max(end, plan_end)
    return end


def last_occupied_date(timeline: dict[str, Any] | None) -> date | None:
    if not timeline:
        return None
    last: date | None = None
    for row in timeline.get("grid", {}).values():
        for d_iso, cell in row.items():
            if cell is None:
                continue
            d = date.fromisoformat(d_iso)
            if last is None or d > last:
                last = d
    return last


def build_markers(
    dates: list[date],
    system_eta: date | None,
    critical_feedback: date | None,
    holidays: dict[str, str] | None = None,
) -> dict[str, list[str]]:
    if holidays is None and dates:
        holidays = _holiday_map_for_span(dates[0], dates[-1])
    holidays = holidays or {}
    markers: dict[str, list[str]] = {}
    for d in dates:
        key = d.isoformat()
        tags: list[str] = []
        if d.weekday() >= 5:
            tags.append("Weekend")
        elif key in holidays:
            # Named Taiwan holiday / makeup day off (weekends already tagged above)
            tags.append("Holiday")
        if system_eta and d == system_eta:
            tags.append("System ETA")
        if critical_feedback and d == critical_feedback:
            tags.append("Critical Feedback")
        markers[key] = tags
    return markers


def blocked_date_set(markers: dict[str, list[str]]) -> set[str]:
    """Weekend / Holiday only (not user Empty marks)."""
    blocked: set[str] = set()
    for d_iso, tags in markers.items():
        if "Weekend" in tags or "Holiday" in tags:
            blocked.add(d_iso)
    return blocked


def calendar_locked_dates(timeline: dict[str, Any]) -> set[str]:
    return blocked_date_set(timeline.get("markers", {}))


def refresh_blocked(timeline: dict[str, Any]) -> None:
    """
    blocked = Weekend/Holiday ∪ user Empty marks.
    Also sync 'Empty' into markers for the Event display row.
    """
    locked = calendar_locked_dates(timeline)
    empty = set(timeline.get("empty_marks", []))
    markers = timeline.setdefault("markers", {})
    for d in timeline.get("dates", []):
        tags = [t for t in markers.get(d, []) if t != "Empty"]
        if d in empty:
            tags.append("Empty")
        markers[d] = tags
    timeline["blocked"] = sorted(locked | empty)
    timeline["empty_marks"] = sorted(empty)
    timeline["calendar_blocked"] = sorted(locked)


def init_timeline_state(
    system_eta: date,
    critical_feedback: date,
    n_systems: int,
    plan_end: date | None = None,
) -> dict[str, Any]:
    """Timeline starts at System ETA; ends at CF+5 business days (or later plan end)."""
    start = system_eta
    probe = max(system_eta, critical_feedback) + timedelta(days=40)
    holidays = _holiday_map_for_span(start, probe)
    end = compute_timeline_end(system_eta, critical_feedback, plan_end, holidays)
    holidays = _holiday_map_for_span(start, end)
    dates = daterange(start, end)
    markers = build_markers(dates, system_eta, critical_feedback, holidays)
    grid: dict[str, dict[str, dict | None]] = {}
    system_order: list[str] = []
    for s in range(1, n_systems + 1):
        row_key = f"System {s}"
        grid[row_key] = {d.isoformat(): None for d in dates}
        system_order.append(row_key)
    tl = {
        "dates": [d.isoformat() for d in dates],
        "markers": markers,
        "empty_marks": [],
        "grid": grid,
        "system_order": system_order,
        "n_systems": n_systems,
        "system_eta": system_eta.isoformat(),
        "critical_feedback": critical_feedback.isoformat(),
        "min_end": add_business_days(critical_feedback, 5, holidays).isoformat(),
    }
    refresh_blocked(tl)
    return tl


def sorted_system_keys(timeline: dict[str, Any]) -> list[str]:
    """Stable system row order (supports custom names, not only 'System N')."""
    grid = timeline.get("grid", {})
    order = timeline.get("system_order") or []
    if order:
        seen: set[str] = set()
        out: list[str] = []
        for k in order:
            if k in grid and k not in seen:
                out.append(k)
                seen.add(k)
        for k in grid:
            if k not in seen:
                out.append(k)
        return out

    def _key(name: str):
        parts = str(name).split()
        if parts and parts[-1].isdigit():
            return (0, int(parts[-1]), name.lower())
        return (1, 0, name.lower())

    return sorted(grid.keys(), key=_key)


_RESERVED_SYSTEM_NAMES = frozenset(
    {
        "",
        "row",
        "empty",
        "marked",
        "event",
        "weekday",
        "date",
        "week",
    }
)


def rename_system_row(
    timeline: dict[str, Any], old_name: str, new_name: str
) -> tuple[dict[str, Any], str | None]:
    """Rename a system row key in grid + system_order. Returns (timeline, error)."""
    old_name = str(old_name or "").strip()
    new_name = str(new_name or "").strip()
    if not old_name:
        return timeline, "Missing system to rename."
    if old_name not in timeline.get("grid", {}):
        return timeline, f"Unknown system row: {old_name}"
    if not new_name:
        return timeline, "System name cannot be empty."
    if new_name.lower() in _RESERVED_SYSTEM_NAMES:
        return timeline, f"'{new_name}' is reserved."
    if new_name == old_name:
        return timeline, None
    if new_name in timeline["grid"]:
        return timeline, f"Name '{new_name}' is already used by another row."

    new_tl = deepcopy(timeline)
    new_tl["grid"][new_name] = new_tl["grid"].pop(old_name)
    order = list(new_tl.get("system_order") or [])
    new_tl["system_order"] = [new_name if k == old_name else k for k in order]
    if new_name not in new_tl["system_order"]:
        # Keep position if old_name was missing from order
        new_tl["system_order"] = sorted_system_keys(new_tl)
    return new_tl, None


def extend_timeline_to(timeline: dict[str, Any], new_end: date) -> dict[str, Any]:
    """Grow calendar columns through new_end (inclusive), preserving grid content."""
    new_tl = deepcopy(timeline)
    dates = [date.fromisoformat(d) for d in new_tl["dates"]]
    if not dates:
        return new_tl
    start = dates[0]
    if new_end <= dates[-1]:
        return new_tl
    system_eta = date.fromisoformat(new_tl["system_eta"]) if new_tl.get("system_eta") else None
    critical_feedback = (
        date.fromisoformat(new_tl["critical_feedback"])
        if new_tl.get("critical_feedback")
        else None
    )
    holidays = _holiday_map_for_span(start, new_end)
    all_dates = daterange(start, new_end)
    markers = build_markers(all_dates, system_eta, critical_feedback, holidays)
    # Preserve Empty tags via empty_marks list
    new_tl["dates"] = [d.isoformat() for d in all_dates]
    new_tl["markers"] = markers
    for row_key, row in new_tl["grid"].items():
        for d in all_dates:
            key = d.isoformat()
            if key not in row:
                row[key] = None
    refresh_blocked(new_tl)
    return new_tl


def ensure_timeline_covers_placement(
    timeline: dict[str, Any],
    last_needed: date,
) -> dict[str, Any]:
    """Extend end to max(CF+5biz, last occupied / needed day)."""
    system_eta = date.fromisoformat(timeline["system_eta"])
    critical_feedback = date.fromisoformat(timeline["critical_feedback"])
    occupied = last_occupied_date(timeline)
    plan_end = last_needed
    if occupied and occupied > plan_end:
        plan_end = occupied
    target = compute_timeline_end(system_eta, critical_feedback, plan_end)
    return extend_timeline_to(timeline, target)


def working_span_dates(
    timeline: dict[str, Any],
    start_date: str,
    duration: int,
) -> tuple[list[str] | None, str]:
    """Collect `duration` fillable dates (skip Weekend / Holiday / Empty marks)."""
    blocked = set(timeline.get("blocked", []))
    if start_date in blocked:
        return None, "Cannot fill Weekend / Holiday / Empty days."
    if start_date not in timeline["dates"]:
        return None, "Start date is outside the timeline."

    collected: list[str] = []
    cur = date.fromisoformat(start_date)
    holidays = _holiday_map_for_span(cur, cur + timedelta(days=duration * 4 + 21))
    empty = set(timeline.get("empty_marks", []))
    guard = 0
    while len(collected) < duration and guard < 500:
        guard += 1
        key = cur.isoformat()
        if is_non_working(cur, holidays) or key in blocked or key in empty:
            cur += timedelta(days=1)
            continue
        collected.append(key)
        cur += timedelta(days=1)
    if len(collected) < duration:
        return None, "Not enough working days to place this item."
    return collected, ""


def _cell_occupied(cell: dict | None) -> bool:
    return cell is not None


def can_place(
    timeline: dict[str, Any],
    row_key: str,
    start_date: str,
    duration: int,
) -> tuple[bool, str, list[str]]:
    blocked = set(timeline.get("blocked", []))
    if start_date in blocked:
        return False, "Cannot fill Weekend / Holiday / Empty days.", []
    span, msg = working_span_dates(timeline, start_date, duration)
    if span is None:
        return False, msg, []
    row = timeline["grid"][row_key]
    for d in span:
        if d in row and _cell_occupied(row.get(d)):
            return False, f"Overlap on {row_key} at {d}.", []
        if d in blocked:
            return False, "Cannot fill Weekend / Holiday / Empty days.", []
    return True, "", span


def place_item(
    timeline: dict[str, Any],
    row_key: str,
    start_date: str,
    test_id: str,
    abbrv: str,
    duration: int,
    is_empty: bool = False,
) -> tuple[dict[str, Any], str | None]:
    """Place a test item (or EMPTY) on working days only. No overlap on same row."""
    ok, msg, span = can_place(timeline, row_key, start_date, duration)
    if not ok:
        return timeline, msg

    new_tl = deepcopy(timeline)
    last_needed = date.fromisoformat(span[-1])
    new_tl = ensure_timeline_covers_placement(new_tl, last_needed)
    refresh_blocked(new_tl)

    # Recompute span after possible extend (blocked may have grown with new weekends)
    ok, msg, span = can_place(new_tl, row_key, start_date, duration)
    if not ok:
        return timeline, msg

    row = new_tl["grid"][row_key]
    for d in span:
        if d not in row:
            for rk in new_tl["grid"]:
                new_tl["grid"][rk].setdefault(d, None)
            row = new_tl["grid"][row_key]
        if _cell_occupied(row.get(d)):
            return timeline, f"Overlap on {row_key} at {d}."

    for offset, d in enumerate(span):
        new_tl["grid"][row_key][d] = {
            "test_id": test_id,
            "abbrv": abbrv if not is_empty else EMPTY_LABEL,
            "duration": duration,
            "offset": offset,
            "is_start": offset == 0,
            "is_empty": is_empty,
            "span_dates": span,
        }
    return new_tl, None


def extract_row_placements(timeline: dict[str, Any], row_key: str) -> list[dict[str, Any]]:
    """Ordered list of started placements (keeps original start + span)."""
    placements: list[dict[str, Any]] = []
    row = timeline["grid"].get(row_key, {})
    for d in timeline["dates"]:
        cell = row.get(d)
        if not cell or not cell.get("is_start"):
            continue
        span = cell.get("span_dates") or [d]
        placements.append(
            {
                "test_id": cell.get("test_id"),
                "abbrv": cell.get("abbrv"),
                "duration": int(cell.get("duration", 1)),
                "is_empty": bool(cell.get("is_empty")),
                "original_start": d,
                "span_dates": list(span),
            }
        )
    return placements


def _first_fillable_on_or_after(timeline: dict[str, Any], day: str) -> str | None:
    """Next fillable day on/after `day` (skip Weekend/Holiday/Empty)."""
    blocked = set(timeline.get("blocked", []))
    dates = timeline.get("dates", [])
    if day not in dates:
        # if day is past end, caller should extend
        if dates and day > dates[-1]:
            return None
        # day before start → use first fillable from start
        for d in dates:
            if d not in blocked:
                return d
        return None
    idx = dates.index(day)
    for d in dates[idx:]:
        if d not in blocked:
            return d
    return None


def _first_fillable_after(timeline: dict[str, Any], day: str) -> str | None:
    """Next fillable day strictly after `day`."""
    dates = timeline.get("dates", [])
    if day in dates:
        idx = dates.index(day) + 1
        if idx >= len(dates):
            return None
        return _first_fillable_on_or_after(timeline, dates[idx])
    if dates and day < dates[0]:
        return _first_fillable_on_or_after(timeline, dates[0])
    return None


def _ensure_fillable_from(
    timeline: dict[str, Any], day: str, *, after: bool = False
) -> tuple[dict[str, Any], str | None]:
    """Extend timeline until a fillable day exists on/after (or strictly after) day."""
    new_tl = timeline
    for _ in range(40):
        refresh_blocked(new_tl)
        start = (
            _first_fillable_after(new_tl, day)
            if after
            else _first_fillable_on_or_after(new_tl, day)
        )
        if start is not None:
            return new_tl, start
        last = date.fromisoformat(new_tl["dates"][-1])
        new_tl = ensure_timeline_covers_placement(new_tl, last + timedelta(days=21))
        refresh_blocked(new_tl)
        for rk in new_tl["grid"]:
            for dd in new_tl["dates"]:
                new_tl["grid"][rk].setdefault(dd, None)
    return new_tl, None


def _place_never_earlier(
    timeline: dict[str, Any],
    row_key: str,
    min_start: str,
    test_id: str,
    abbrv: str,
    duration: int,
    is_empty: bool,
) -> tuple[dict[str, Any], str | None, str | None]:
    """
    Place item starting at first fillable on/after min_start.
    Returns (timeline, error, last_span_date).
    """
    new_tl = timeline
    cursor = min_start
    for _ in range(80):
        new_tl, start = _ensure_fillable_from(new_tl, cursor, after=False)
        if start is None:
            return new_tl, f"{row_key}: no fillable day for {test_id}", None
        trial, err = place_item(
            new_tl, row_key, start, test_id, abbrv, duration, is_empty
        )
        if err:
            # bump past this start and retry
            idx = new_tl["dates"].index(start) if start in new_tl["dates"] else -1
            if idx < 0 or idx + 1 >= len(new_tl["dates"]):
                last = date.fromisoformat(new_tl["dates"][-1])
                new_tl = ensure_timeline_covers_placement(new_tl, last + timedelta(days=14))
                refresh_blocked(new_tl)
                for rk in new_tl["grid"]:
                    for dd in new_tl["dates"]:
                        new_tl["grid"][rk].setdefault(dd, None)
                cursor = new_tl["dates"][-1]
            else:
                cursor = new_tl["dates"][idx + 1]
            continue
        cell = trial["grid"][row_key].get(start) or {}
        span = cell.get("span_dates") or [start]
        return trial, None, span[-1]
    return new_tl, f"{row_key}: failed to place {test_id}", None


def postpone_system_schedules(timeline: dict[str, Any]) -> tuple[dict[str, Any], list[str]]:
    """
    After Empty marks change, re-place each system item so that:
      - It never starts earlier than its original start (blank days before stay blank)
      - Full original duration is kept as ONE item
      - Empty / Weekend / Holiday days are skipped inside the span
        e.g. start 09/30, Empty on 10/01, duration 3 → 09/30, 10/02, 10/05
    """
    new_tl = deepcopy(timeline)
    refresh_blocked(new_tl)
    errors: list[str] = []

    for row_key in sorted_system_keys(new_tl):
        placements = extract_row_placements(new_tl, row_key)
        for d in list(new_tl["grid"][row_key].keys()):
            new_tl["grid"][row_key][d] = None

        # Floor for non-overlap with prior items on this row only — NOT timeline start.
        # None means "no extra floor" so a blank day before the first item stays blank.
        earliest_free: str | None = None

        for p in placements:
            min_start = p["original_start"]
            if earliest_free and earliest_free > min_start:
                min_start = earliest_free

            new_tl, err, last = _place_never_earlier(
                new_tl,
                row_key,
                min_start,
                p["test_id"],
                p["abbrv"],
                p["duration"],  # keep full duration; place_item skips Empty/weekend
                p["is_empty"],
            )
            if err:
                errors.append(err)
                continue
            if last and last in new_tl["dates"]:
                li = new_tl["dates"].index(last)
                earliest_free = (
                    new_tl["dates"][li + 1]
                    if li + 1 < len(new_tl["dates"])
                    else last
                )

    refresh_blocked(new_tl)
    return new_tl, errors


def shift_fillable_day(
    timeline: dict[str, Any],
    day: str,
    delta: int,
) -> tuple[dict[str, Any], str | None]:
    """
    Move `day` by `delta` fillable days (+ postpone / − ahead).
    Extends timeline when postponing past the end. Returns (timeline, new_day|None).
    """
    if delta == 0:
        return timeline, day
    new_tl = deepcopy(timeline)
    refresh_blocked(new_tl)
    blocked = set(new_tl.get("blocked", []))
    dates = new_tl.get("dates", [])
    if day not in dates:
        return new_tl, None
    idx = dates.index(day)
    step = 1 if delta > 0 else -1
    need = abs(delta)
    found = 0
    i = idx
    guard = 0
    while found < need and guard < 500:
        guard += 1
        i += step
        if i < 0:
            return new_tl, None
        if i >= len(new_tl["dates"]):
            last = date.fromisoformat(new_tl["dates"][-1])
            new_tl = ensure_timeline_covers_placement(new_tl, last + timedelta(days=21))
            refresh_blocked(new_tl)
            blocked = set(new_tl.get("blocked", []))
            for rk in new_tl["grid"]:
                for dd in new_tl["dates"]:
                    new_tl["grid"][rk].setdefault(dd, None)
            continue
        if new_tl["dates"][i] not in blocked:
            found += 1
            if found == need:
                return new_tl, new_tl["dates"][i]
    return new_tl, None


def shift_system_schedule(
    timeline: dict[str, Any],
    row_key: str,
    delta: int,
) -> tuple[dict[str, Any], list[str]]:
    """
    Shift every test item on one system by `delta` fillable days
    (+1 postpone, −1 ahead). Blank days before items stay blank when possible.
    """
    if row_key not in timeline.get("grid", {}):
        return timeline, [f"Unknown system: {row_key}"]
    if delta == 0:
        return timeline, []

    new_tl = deepcopy(timeline)
    refresh_blocked(new_tl)
    placements = extract_row_placements(new_tl, row_key)
    if not placements:
        return new_tl, []

    for d in list(new_tl["grid"][row_key].keys()):
        new_tl["grid"][row_key][d] = None

    moved: list[tuple[str, dict[str, Any]]] = []
    errors: list[str] = []
    for p in placements:
        new_tl, new_start = shift_fillable_day(new_tl, p["original_start"], delta)
        if new_start is None:
            if delta < 0:
                # cannot move earlier — keep original start
                new_start = p["original_start"]
                errors.append(
                    f"{row_key}: {p['abbrv']} already at earliest fillable day"
                )
            else:
                errors.append(f"{row_key}: could not postpone {p['abbrv']}")
                continue
        moved.append((new_start, p))

    moved.sort(key=lambda x: x[0])
    earliest_free: str | None = None
    for new_start, p in moved:
        min_start = new_start
        if earliest_free and earliest_free > min_start:
            min_start = earliest_free
        new_tl, err, last = _place_never_earlier(
            new_tl,
            row_key,
            min_start,
            p["test_id"],
            p["abbrv"],
            p["duration"],
            p["is_empty"],
        )
        if err:
            errors.append(err)
            continue
        if last and last in new_tl["dates"]:
            li = new_tl["dates"].index(last)
            earliest_free = (
                new_tl["dates"][li + 1]
                if li + 1 < len(new_tl["dates"])
                else last
            )

    refresh_blocked(new_tl)
    return new_tl, errors


def shift_system_from_date(
    timeline: dict[str, Any],
    row_key: str,
    from_date: str,
    delta: int,
) -> tuple[dict[str, Any], list[str]]:
    """
    Shift only test items on `row_key` whose span starts on/after `from_date`
    by `delta` fillable days (+1 postpone, −1 ahead).

    If the selected cell is mid-span of an item, that whole item and everything
    after it move. Items strictly before the cutoff stay fixed.

    On ahead: if any moved item would overlap a kept item (or fail to place),
    abort with no changes and return a warning.
    """
    if row_key not in timeline.get("grid", {}):
        return timeline, [f"Unknown system: {row_key}"]
    if delta == 0:
        return timeline, []
    if from_date not in timeline.get("dates", []):
        return timeline, [f"Date {from_date} is outside the timeline."]

    placements = extract_row_placements(timeline, row_key)
    if not placements:
        return timeline, [f"{row_key}: no test items to shift."]

    # If user selected a day inside an item, cutoff = that item's start
    cutoff = from_date
    cell = timeline["grid"][row_key].get(from_date)
    if cell and not cell.get("is_empty"):
        for p in placements:
            if from_date == p["original_start"] or from_date in (p.get("span_dates") or []):
                cutoff = p["original_start"]
                break

    keep = [p for p in placements if p["original_start"] < cutoff]
    move = [p for p in placements if p["original_start"] >= cutoff]
    if not move:
        return timeline, [
            f"{row_key}: no test items on/after {format_date_display(from_date)} to shift."
        ]

    # Work on a copy; clear only the items that will move
    new_tl = deepcopy(timeline)
    refresh_blocked(new_tl)
    for p in move:
        for d in p.get("span_dates") or [p["original_start"]]:
            if d in new_tl["grid"][row_key]:
                new_tl["grid"][row_key][d] = None

    proposed: list[tuple[str, dict[str, Any]]] = []
    for p in move:
        new_tl, new_start = shift_fillable_day(new_tl, p["original_start"], delta)
        if new_start is None:
            action = "ahead" if delta < 0 else "postpone"
            return timeline, [
                f"Cannot {action} from {format_date_display(from_date)}: "
                f"{p['abbrv']} has no fillable day to move to."
            ]
        proposed.append((new_start, p))

    proposed.sort(key=lambda x: x[0])

    # Place all proposed moves; any overlap → full abort (especially ahead into kept items)
    trial = deepcopy(new_tl)
    for new_start, p in proposed:
        ok, msg, _span = can_place(trial, row_key, new_start, int(p["duration"]))
        if not ok:
            action = "ahead" if delta < 0 else "postpone"
            return timeline, [
                f"Cannot {action} from {format_date_display(from_date)}: "
                f"{p['abbrv']} would overlap ({msg}). Schedule unchanged."
            ]
        trial, err = place_item(
            trial,
            row_key,
            new_start,
            p["test_id"],
            p["abbrv"],
            int(p["duration"]),
            bool(p["is_empty"]),
        )
        if err:
            action = "ahead" if delta < 0 else "postpone"
            return timeline, [
                f"Cannot {action} from {format_date_display(from_date)}: "
                f"{p['abbrv']} — {err} Schedule unchanged."
            ]

    refresh_blocked(trial)
    return trial, []


def refresh_durations_from_catalog(
    timeline: dict[str, Any],
    info: dict[str, dict],
) -> tuple[dict[str, Any], list[str]]:
    """
    Re-apply each placed test item using Duration_Days from `info` (loaded catalog).

    Keeps the current schedule sequence: same items, same original start dates
    (and blank gaps). Does NOT reload the Case sequence template.

    If a longer duration would collide with a later item, that later item is
    pushed forward (never earlier than its original start).
    """
    new_tl = deepcopy(timeline)
    refresh_blocked(new_tl)
    errors: list[str] = []

    for row_key in sorted_system_keys(new_tl):
        placements = extract_row_placements(new_tl, row_key)
        if not placements:
            continue

        for d in list(new_tl["grid"][row_key].keys()):
            new_tl["grid"][row_key][d] = None

        earliest_free: str | None = None
        for p in placements:
            tid = str(p.get("test_id") or "")
            is_empty = bool(p.get("is_empty"))
            if is_empty:
                dur = 1
                abbrv = p.get("abbrv") or EMPTY_LABEL
            else:
                meta = info.get(tid)
                if meta:
                    dur = int(meta["Duration_Days"])
                    abbrv = meta.get("Abbrv_Name") or p.get("abbrv") or tid
                else:
                    dur = int(p["duration"])
                    abbrv = p.get("abbrv") or tid
                    errors.append(
                        f"{row_key}: {tid} not in loaded data — kept previous duration."
                    )

            min_start = p["original_start"]
            if earliest_free and earliest_free > min_start:
                min_start = earliest_free

            new_tl, err, last = _place_never_earlier(
                new_tl,
                row_key,
                min_start,
                tid,
                abbrv,
                dur,
                is_empty,
            )
            if err:
                errors.append(err)
                continue
            if last and last in new_tl["dates"]:
                li = new_tl["dates"].index(last)
                earliest_free = (
                    new_tl["dates"][li + 1]
                    if li + 1 < len(new_tl["dates"])
                    else last
                )

    refresh_blocked(new_tl)
    return new_tl, errors


def set_calendar_empty_mark(
    timeline: dict[str, Any],
    day: str,
    enabled: bool,
) -> tuple[dict[str, Any], list[str]]:
    """
    Toggle Empty on the Event row for a day.
    When enabling Empty: that day cannot be filled; existing plans keep their
    original start (never earlier) and skip Empty days while keeping duration.
    """
    new_tl = deepcopy(timeline)
    locked = calendar_locked_dates(new_tl)
    if day in locked:
        return timeline, [f"{day} is already Weekend/Holiday."]
    if day not in new_tl.get("dates", []):
        return timeline, [f"{day} is outside the timeline."]

    empty = set(new_tl.get("empty_marks", []))
    if enabled:
        empty.add(day)
    else:
        empty.discard(day)
    new_tl["empty_marks"] = sorted(empty)
    refresh_blocked(new_tl)

    if enabled:
        new_tl, errors = postpone_system_schedules(new_tl)
        return new_tl, errors
    return new_tl, []


def clear_cell_span(timeline: dict[str, Any], row_key: str, date_iso: str) -> dict[str, Any]:
    """Clear the whole span belonging to the cell at date_iso."""
    new_tl = deepcopy(timeline)
    cell = new_tl["grid"][row_key].get(date_iso)
    if not cell:
        return new_tl
    span = cell.get("span_dates")
    if span:
        for d in span:
            if d in new_tl["grid"][row_key]:
                new_tl["grid"][row_key][d] = None
        return new_tl
    # Fallback for legacy contiguous cells
    dates = new_tl["dates"]
    if date_iso not in dates:
        return new_tl
    idx = dates.index(date_iso)
    start_idx = idx - int(cell.get("offset", 0))
    duration = int(cell.get("duration", 1))
    for offset in range(duration):
        d = dates[start_idx + offset]
        new_tl["grid"][row_key][d] = None
    return new_tl


def display_label(cell: dict | None) -> str:
    if not cell:
        return ""
    if cell.get("is_empty"):
        return EMPTY_LABEL if cell.get("is_start") else "…"
    if cell.get("is_start"):
        return str(cell.get("abbrv") or cell.get("test_id"))
    return "…"


def editor_cell_label(cell: dict | None) -> str:
    """Value shown in Excel-like selectbox cells (must match select options)."""
    if not cell:
        return ""
    if cell.get("is_empty"):
        return EMPTY_LABEL
    tid = cell.get("test_id", "")
    abbrv = cell.get("abbrv") or tid
    return f"{abbrv} ({tid})"


def timeline_calendar_df(timeline: dict[str, Any]) -> pd.DataFrame:
    """
    Single merged calendar table (Excel-like):
      Date | Weekday | Event | System 1 | System 2 | ...
    """
    dates = timeline["dates"]
    markers = timeline["markers"]
    grid = timeline["grid"]
    system_keys = sorted_system_keys(timeline)

    rows = []
    for d in dates:
        row = {
            "Date": d,
            "Weekday": date.fromisoformat(d).strftime("%a"),
            "Event": " | ".join(markers.get(d, [])),
        }
        for sk in system_keys:
            row[sk] = editor_cell_label(grid[sk].get(d))
        rows.append(row)
    return pd.DataFrame(rows)


def style_calendar_df(df: pd.DataFrame, timeline: dict[str, Any]):
    """Light-red entire Weekend / Holiday rows."""
    blocked = set(timeline.get("blocked", []))

    def _row_style(row: pd.Series):
        if str(row.get("Date", "")) in blocked:
            return [f"background-color: {BLOCKED_FILL}"] * len(row)
        return [""] * len(row)

    return df.style.apply(_row_style, axis=1)


# Keep aliases used by export preview styling of wide format
def timeline_header_df(timeline: dict[str, Any]) -> pd.DataFrame:
    """Read-only Weekday + Event rows (wide format)."""
    dates = timeline["dates"]
    markers = timeline["markers"]
    weekdays = [date.fromisoformat(d).strftime("%a") for d in dates]
    marks = [" | ".join(markers.get(d, [])) for d in dates]
    return pd.DataFrame(
        [
            {"Row": "Weekday", **{d: weekdays[i] for i, d in enumerate(dates)}},
            {"Row": "Event", **{d: marks[i] for i, d in enumerate(dates)}},
        ]
    )


def timeline_editor_df(timeline: dict[str, Any]) -> pd.DataFrame:
    """Editable System rows for data_editor selectboxes (wide format)."""
    dates = timeline["dates"]
    grid = timeline["grid"]
    rows = []
    for row_key in sorted_system_keys(timeline):
        row = {"Row": row_key}
        for d in dates:
            row[d] = editor_cell_label(grid[row_key].get(d))
        rows.append(row)
    return pd.DataFrame(rows)


def timeline_to_export_df(
    timeline: dict[str, Any],
    *,
    weight_kg: float | None = None,
    detail_by_id: dict[str, dict[str, str]] | None = None,
) -> pd.DataFrame:
    """
    Export layout:
      Row 0: Date
      Row 1: Event
      Then for each system:
        - schedule row (test items)
        - profile row under it (weight-band notes; blank if none)
      First-column label is the system name on the schedule row and blank on the
      profile row so Excel can vertically merge them into one cell.
    """
    dates = timeline["dates"]
    markers = timeline["markers"]
    grid = timeline["grid"]

    profile_col = None
    if weight_kg is not None and detail_by_id is not None:
        profile_col = profile_column_for_weight(float(weight_kg))

    rows: list[dict[str, str]] = []
    date_row = {"Row": "Date"}
    mark_row = {"Row": "Event"}
    for d in dates:
        date_row[d] = format_date_display(d)
        mark_row[d] = " | ".join(markers.get(d, []))
    rows.append(date_row)
    rows.append(mark_row)

    for row_key in sorted_system_keys(timeline):
        schedule = {"Row": row_key}
        profile = {"Row": ""}  # blank — merged with schedule label in Excel
        for d in dates:
            cell = grid[row_key].get(d)
            if not cell:
                schedule[d] = ""
                profile[d] = ""
            elif cell.get("is_empty"):
                schedule[d] = EMPTY_TOKEN
                profile[d] = ""
            else:
                tid = str(cell.get("test_id", "") or "").strip()
                abbrv = (cell.get("abbrv") or tid).strip()
                schedule[d] = f"{abbrv} ({tid})" if tid else abbrv
                note = ""
                if profile_col and detail_by_id and tid:
                    meta = detail_by_id.get(tid)
                    if meta:
                        note = (meta.get(profile_col, "") or "").strip()
                profile[d] = note
        rows.append(schedule)
        rows.append(profile)

    return pd.DataFrame(rows)


def timeline_to_display_df(timeline: dict[str, Any]) -> pd.DataFrame:
    """Human-readable grid: dates as columns, marker + systems as rows."""
    dates = timeline["dates"]
    markers = timeline["markers"]
    grid = timeline["grid"]

    weekdays = []
    marks = []
    for d in dates:
        dt = date.fromisoformat(d)
        weekdays.append(dt.strftime("%a"))
        marks.append(" | ".join(markers.get(d, [])))
    rows = [
        {"Row": "Weekday", **{d: weekdays[i] for i, d in enumerate(dates)}},
        {"Row": "Event", **{d: marks[i] for i, d in enumerate(dates)}},
    ]
    for row_key in sorted_system_keys(timeline):
        row = {"Row": row_key}
        for d in dates:
            row[d] = display_label(grid[row_key].get(d))
        rows.append(row)
    return pd.DataFrame(rows)


def style_timeline_display(df: pd.DataFrame, timeline: dict[str, Any]):
    """Light-red Weekend/Holiday columns; light-yellow System ETA / Critical Feedback on Event row."""
    blocked = set(timeline.get("blocked", []))
    markers = timeline.get("markers", {})
    milestone_dates = {
        d
        for d, tags in markers.items()
        if "System ETA" in tags or "Critical Feedback" in tags
    }

    def _row_style(row: pd.Series):
        is_event = str(row.get("Row", "")) == "Event"
        styles: list[str] = []
        for col in row.index:
            if col == "Row":
                styles.append("")
            elif col in blocked:
                styles.append(f"background-color: {BLOCKED_FILL}")
            elif is_event and col in milestone_dates:
                styles.append(f"background-color: {EVENT_MILESTONE_FILL}")
            else:
                styles.append("")
        return styles

    return df.style.apply(_row_style, axis=1)


def export_timeline_xlsx(df: pd.DataFrame, timeline: dict[str, Any]) -> bytes:
    """
    Write export table to XLSX with the same fills as the on-screen table:
      - Weekend / Holiday columns → light red
      - Event row System ETA / Critical Feedback → light yellow
      - System name + profile rows: first column vertically merged
    CSV cannot store colors or merges; use this for colored downloads.
    """
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    blocked = set(timeline.get("blocked", []))
    markers = timeline.get("markers", {})
    milestone_dates = {
        d
        for d, tags in markers.items()
        if "System ETA" in tags or "Critical Feedback" in tags
    }
    fill_blocked = PatternFill(
        start_color=BLOCKED_FILL.lstrip("#"),
        end_color=BLOCKED_FILL.lstrip("#"),
        fill_type="solid",
    )
    fill_milestone = PatternFill(
        start_color=EVENT_MILESTONE_FILL.lstrip("#"),
        end_color=EVENT_MILESTONE_FILL.lstrip("#"),
        fill_type="solid",
    )
    header_font = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    center_left = Alignment(wrap_text=True, vertical="center", horizontal="left")

    wb = Workbook()
    ws = wb.active
    ws.title = "Timeline"

    columns = list(df.columns)
    for col_idx, col_name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=str(col_name))
        cell.font = header_font
        if col_name in blocked:
            cell.fill = fill_blocked

    records = df.to_dict(orient="records")
    excel_row = 2
    i = 0
    while i < len(records):
        record = records[i]
        row_label = "" if record.get("Row") is None else str(record.get("Row"))
        next_label = None
        if i + 1 < len(records):
            nxt = records[i + 1].get("Row")
            next_label = "" if nxt is None else str(nxt)

        # System schedule + following blank-label profile row → merge first column
        merge_pair = (
            row_label not in ("", "Date", "Event")
            and next_label == ""
        )

        def _write_data_row(rec: dict, r_idx: int, *, is_event: bool) -> None:
            for col_idx, col_name in enumerate(columns, start=1):
                raw = rec.get(col_name, "")
                if raw is None or (isinstance(raw, float) and pd.isna(raw)):
                    value = ""
                else:
                    value = str(raw)
                cell = ws.cell(row=r_idx, column=col_idx, value=value)
                cell.alignment = wrap
                if col_name == "Row":
                    continue
                if col_name in blocked:
                    cell.fill = fill_blocked
                elif is_event and col_name in milestone_dates:
                    cell.fill = fill_milestone

        if merge_pair:
            _write_data_row(record, excel_row, is_event=False)
            _write_data_row(records[i + 1], excel_row + 1, is_event=False)
            # Merge first column; keep system name centered across both rows
            ws.merge_cells(
                start_row=excel_row,
                start_column=1,
                end_row=excel_row + 1,
                end_column=1,
            )
            label_cell = ws.cell(row=excel_row, column=1, value=row_label)
            label_cell.alignment = center_left
            excel_row += 2
            i += 2
        else:
            _write_data_row(record, excel_row, is_event=(row_label == "Event"))
            excel_row += 1
            i += 1

    ws.column_dimensions["A"].width = 16
    for col_idx in range(2, len(columns) + 1):
        letter = ws.cell(row=1, column=col_idx).column_letter
        ws.column_dimensions[letter].width = 18

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def fillable_dates(timeline: dict[str, Any]) -> list[str]:
    blocked = set(timeline.get("blocked", []))
    return [d for d in timeline["dates"] if d not in blocked]


def apply_sequence_template(
    timeline: dict[str, Any],
    sequence_df: pd.DataFrame,
    start_from_date: str | None = None,
    *,
    info: dict[str, dict] | None = None,
    account: str | None = None,
) -> tuple[dict[str, Any], list[str]]:
    """
    Pre-fill timeline from an All-Case sequence sheet.

    Preferred format — one token per cell (columns 1, 2, 3, …):
      Row | 1 | 2 | 3 | 4 | 5 | 6
      DUT-A | 1 | T003 | T004 | T006 | 5 | T010
    Column A may be any system label (not only "System 1"). Those names are used on the timeline.
    Integers = blank fillable days; Test_ID = place that item (Duration_Days from catalog).
    """
    if info is None:
        if not account:
            raise ValueError("Provide info or account for apply_sequence_template")
        info = test_info_by_id(account=account)
    errors: list[str] = []
    new_tl = deepcopy(timeline)
    fillable = fillable_dates(new_tl)
    if not fillable:
        return new_tl, ["Timeline has no fillable (working) days."]

    start_idx = 0
    if start_from_date:
        for i, d in enumerate(fillable):
            if d >= start_from_date:
                start_idx = i
                break

    row_col = "Row" if "Row" in sequence_df.columns else sequence_df.columns[0]

    def _ensure_cursor(tl: dict[str, Any], cursor: int, fill: list[str]) -> tuple[dict[str, Any], list[str]]:
        while cursor >= len(fill):
            last = date.fromisoformat(tl["dates"][-1])
            tl = ensure_timeline_covers_placement(tl, last + timedelta(days=14))
            fill = fillable_dates(tl)
            if cursor >= len(fill):
                tl = extend_timeline_to(tl, last + timedelta(days=30))
                fill = fillable_dates(tl)
                if cursor >= len(fill):
                    break
        return tl, fill

    def _apply_tokens(row_name: str, tokens: list[str], cursor: int) -> int:
        nonlocal new_tl, fillable, errors
        for token in tokens:
            new_tl, fillable = _ensure_cursor(new_tl, cursor, fillable)
            if cursor >= len(fillable):
                errors.append(f"Ran out of dates for {row_name}.")
                break

            if token.isdigit() or (token.startswith("-") and token[1:].isdigit()):
                blank_days = int(token)
                if blank_days < 0:
                    errors.append(f"{row_name}: negative blank count '{token}' ignored.")
                    continue
                cursor += blank_days
                new_tl, fillable = _ensure_cursor(new_tl, cursor, fillable)
                continue

            tid = token.strip()
            if tid.upper() == EMPTY_TOKEN or tid.lower() == "empty":
                place_date = fillable[cursor]
                new_tl, err = place_item(
                    new_tl, row_name, place_date, EMPTY_TOKEN, EMPTY_LABEL, 1, True
                )
                if err:
                    errors.append(err)
                fillable = fillable_dates(new_tl)
                cursor += 1
                continue

            meta = info.get(tid)
            if not meta:
                errors.append(f"Unknown Test_ID '{tid}' in sequence for {row_name}.")
                continue
            place_date = fillable[cursor]
            dur = int(meta["Duration_Days"])
            new_tl, fillable = _ensure_cursor(new_tl, cursor + dur - 1, fillable)
            new_tl, err = place_item(
                new_tl,
                row_name,
                place_date,
                meta["Test_ID"],
                meta["Abbrv_Name"],
                dur,
                False,
            )
            if err:
                errors.append(f"{row_name} @ {place_date}: {err}")
                cursor += 1
            else:
                fillable = fillable_dates(new_tl)
                cursor += dur
        return cursor

    # Collect sequence rows — first column may be any label (not only "System N")
    skip_names = {
        "",
        "row",
        "empty",
        "marked",
        "event",
        "weekday",
        "date",
        "week",
    }
    seq_entries: list[tuple[str, list[str], Any]] = []
    for _, seq_row in sequence_df.iterrows():
        row_name = str(seq_row.get(row_col, "")).strip()
        if not row_name or row_name.lower() in skip_names:
            continue
        tokens = tokens_from_sequence_row(seq_row, sequence_df.columns)
        seq_entries.append((row_name, tokens, seq_row))

    if not seq_entries:
        return new_tl, errors

    # Remap timeline rows by order: sequence row 1 → first system, etc.
    # Custom names from the sheet replace default "System 1", "System 2", …
    old_keys = sorted_system_keys(new_tl)
    dates = list(new_tl["dates"])
    new_grid: dict[str, dict] = {}
    new_order: list[str] = []
    used_names: set[str] = set()

    for i, old_key in enumerate(old_keys):
        if i < len(seq_entries):
            name = seq_entries[i][0]
            # Avoid duplicate keys
            base = name
            n = 2
            while name in used_names:
                name = f"{base} ({n})"
                n += 1
        else:
            name = old_key
            while name in used_names:
                name = f"{name}_"
        used_names.add(name)
        new_grid[name] = {d: None for d in dates}
        new_order.append(name)

    new_tl["grid"] = new_grid
    new_tl["system_order"] = new_order
    new_tl["n_systems"] = len(new_order)

    for i, (row_name, tokens, seq_row) in enumerate(seq_entries):
        if i >= len(new_order):
            break
        target = new_order[i]
        cursor = start_idx
        if tokens:
            _apply_tokens(target, tokens, cursor)
            continue

        # Legacy Day_1 / Day_2 … calendar-slot layout
        day_cols = [c for c in sequence_df.columns if str(c).startswith("Day_")]
        if not day_cols:
            continue
        for col in day_cols:
            new_tl, fillable = _ensure_cursor(new_tl, cursor, fillable)
            if cursor >= len(fillable):
                errors.append(f"Ran out of dates for {target}.")
                break
            raw = seq_row.get(col)
            if pd.isna(raw) or str(raw).strip() == "":
                cursor += 1
                continue
            token = str(raw).strip()
            place_date = fillable[cursor]
            if token.upper() == EMPTY_TOKEN or token.lower() == "empty":
                new_tl, err = place_item(
                    new_tl, target, place_date, EMPTY_TOKEN, EMPTY_LABEL, 1, True
                )
                if err:
                    errors.append(err)
                fillable = fillable_dates(new_tl)
                cursor += 1
                continue
            meta = info.get(token)
            if not meta:
                errors.append(f"Unknown Test_ID '{token}' in sequence for {target}.")
                cursor += 1
                continue
            dur = int(meta["Duration_Days"])
            new_tl, err = place_item(
                new_tl,
                target,
                place_date,
                meta["Test_ID"],
                meta["Abbrv_Name"],
                dur,
                False,
            )
            if err:
                errors.append(f"{target} @ {place_date}: {err}")
                cursor += 1
            else:
                fillable = fillable_dates(new_tl)
                cursor += dur
    return new_tl, errors
