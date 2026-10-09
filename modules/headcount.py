from __future__ import annotations

from io import BytesIO
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import get_column_letter

from modules.config import (
    HC_CALENDAR_MONTHS,
    HC_DIST_PER_MONTH,
    HC_DIST_PER_PHASE,
    HC_DISTRIBUTION_OPTS,
    HC_ENGINEER_TYPES,
    HC_SHOW_MONTH,
    HC_SHOW_SEQUENCE,
    PHASES,
)

# Re-export distribution labels for callers
__all__ = [
    "HC_DIST_PER_MONTH",
    "HC_DIST_PER_PHASE",
    "HC_DISTRIBUTION_OPTS",
    "HC_SHOW_MONTH",
    "HC_SHOW_SEQUENCE",
    "allocate_task_headcount",
    "build_headcount_table",
    "build_headcount_detail_table",
    "compute_engineer_month_values",
    "calendar_month_labels",
    "export_headcount_xlsx",
]


def calendar_month_labels(
    n_months: int, *, project_start_month: str = "Jan"
) -> list[str]:
    """
    Calendar month names for a timeline of n_months starting at project_start_month.
    Wraps Dec → Jan as needed (e.g. Jun → Jul → … → Dec → Jan → …).
    """
    n = max(int(n_months), 0)
    if n <= 0:
        return []
    names = list(HC_CALENDAR_MONTHS)
    start = str(project_start_month or "Jan").strip()
    try:
        idx = names.index(start)
    except ValueError:
        idx = 0
    return [names[(idx + i) % 12] for i in range(n)]


def _round_headcount(value: float) -> float:
    """Round a monthly headcount total to 1 decimal (standard rounding)."""
    return round(float(value), 1)


def allocate_task_headcount(
    headcount: float,
    n_months: int,
    distribution: str,
    *,
    start_month: int | None = None,
    end_month: int | None = None,
) -> list[float]:
    """
    Spread a task headcount across months in one phase.

    Per Month  → add the full value to every month in the phase.
    Per Phase  → divide by the selected month span and put the share only
                 in months start_month…end_month (1-based within the phase).
                 Defaults to the full phase (1…n_months) when omitted.
    """
    n = max(int(n_months), 0)
    if n <= 0:
        return []
    value = float(headcount or 0)
    mode = str(distribution or HC_DIST_PER_MONTH).strip()
    if mode == HC_DIST_PER_PHASE:
        start = int(start_month) if start_month is not None else 1
        end = int(end_month) if end_month is not None else n
        start = max(1, min(start, n))
        end = max(start, min(end, n))
        span = end - start + 1
        share = value / span if span > 0 else 0.0
        out = [0.0] * n
        for m in range(start, end + 1):
            out[m - 1] = share
        return out
    # Default: Per Month
    return [value] * n


def _blank_month_values(n: int) -> list[float]:
    return [0.0] * n


def compute_engineer_month_values(
    phase_cells: list[str],
    *,
    months_by_phase: dict[str, Any],
    tasks_by_phase: dict[str, dict[str, list[dict[str, Any]]]],
) -> dict[str, list[Any]]:
    """
    Sum task allocations into per-month values for each engineer row.

    tasks_by_phase[phase][engineer] = [
        {
            "task": str,
            "headcount": float,
            "distribution": "Per Month"|"Per Phase",
            "start_month": int | None,  # 1-based within phase (Per Phase)
            "end_month": int | None,
        },
        ...
    ]
    """
    n = len(phase_cells)
    totals = {eng: _blank_month_values(n) for eng in HC_ENGINEER_TYPES}
    if n == 0:
        return {eng: [] for eng in HC_ENGINEER_TYPES}

    # Index ranges per phase within the flat month timeline
    phase_ranges: dict[str, tuple[int, int]] = {}
    i = 0
    while i < n:
        phase = phase_cells[i]
        start = i
        while i < n and phase_cells[i] == phase:
            i += 1
        phase_ranges[phase] = (start, i)

    for phase, (start, end) in phase_ranges.items():
        n_months = end - start
        eng_tasks = (tasks_by_phase or {}).get(phase) or {}
        for eng in HC_ENGINEER_TYPES:
            for task in eng_tasks.get(eng) or []:
                hc = float(task.get("headcount") or 0)
                if hc == 0:
                    continue
                dist = str(task.get("distribution") or HC_DIST_PER_MONTH)
                start_m = task.get("start_month")
                end_m = task.get("end_month")
                if dist != HC_DIST_PER_PHASE:
                    start_m = None
                    end_m = None
                shares = allocate_task_headcount(
                    hc,
                    n_months,
                    dist,
                    start_month=start_m,
                    end_month=end_m,
                )
                for offset, share in enumerate(shares):
                    if abs(share) > 1e-15:
                        totals[eng][start + offset] += share

    # Round monthly totals to 1 decimal; keep 0 (not None) when HC is zero.
    out: dict[str, list[Any]] = {}
    for eng in HC_ENGINEER_TYPES:
        out[eng] = [_round_headcount(v) for v in totals[eng]]
    return out


def build_headcount_table(
    phases: list[str],
    months_by_phase: dict[str, Any],
    *,
    tasks_by_phase: dict[str, dict[str, list[dict[str, Any]]]] | None = None,
    engineer_values: dict[str, list[Any]] | None = None,
    row_comments: dict[str, Any] | None = None,
    month_display: str = HC_SHOW_SEQUENCE,
    project_start_month: str = "Jan",
) -> pd.DataFrame:
    """
    Build a headcount grid matching HEADCOUNT_TEMPLATE_*.xlsx:

      Row 0: Phase   | <phase repeated for each of its months> | Comments
      Row 1: Month   | 1, 2, 3, …  OR  Jun, Jul, Aug, … | Comments
      Row 2–4: Structure / CAE / S&V Engineer monthly values | Comments

    month_display:
      Show in sequence → Month row is 1, 2, 3, …
      Show in month    → Month row is calendar names from project_start_month
    """
    ordered = [p for p in PHASES if p in {str(x) for x in phases}]
    if not ordered:
        ordered = [str(p) for p in phases]

    phase_cells: list[str] = []
    month_seq: list[int] = []
    month_num = 1
    for phase in ordered:
        n = int(months_by_phase.get(phase, 0) or 0)
        if n <= 0:
            continue
        for _ in range(n):
            phase_cells.append(phase)
            month_seq.append(month_num)
            month_num += 1

    label_col = ""
    comments_col = "Comments"
    row_labels = ["Phase", "Month", *HC_ENGINEER_TYPES]
    comments = row_comments or {}

    def _comment_for(label: str) -> Any:
        val = comments.get(label)
        if val is None or (isinstance(val, float) and pd.isna(val)):
            return None
        text = str(val).strip()
        return text or None

    if not month_seq:
        return pd.DataFrame(
            {
                label_col: row_labels,
                comments_col: [_comment_for(lbl) for lbl in row_labels],
            }
        )

    use_calendar = str(month_display or "").strip() == HC_SHOW_MONTH
    if use_calendar:
        month_labels: list[Any] = calendar_month_labels(
            len(month_seq), project_start_month=project_start_month
        )
        # Unique column keys when the same month name repeats across years
        col_keys: list[str] = []
        seen: dict[str, int] = {}
        for label in month_labels:
            key = str(label)
            if key in seen:
                seen[key] += 1
                col_keys.append(f"{key}_{seen[key]}")
            else:
                seen[key] = 1
                col_keys.append(key)
    else:
        month_labels = list(month_seq)
        col_keys = [str(m) for m in month_seq]

    if engineer_values is not None:
        eng = {name: list(vals) for name, vals in engineer_values.items()}
        for name in HC_ENGINEER_TYPES:
            eng.setdefault(name, [None] * len(month_seq))
    else:
        eng = compute_engineer_month_values(
            phase_cells,
            months_by_phase=months_by_phase,
            tasks_by_phase=tasks_by_phase or {},
        )

    def _eng_row(name: str) -> list[Any]:
        vals = list(eng.get(name) or [])
        out: list[Any] = []
        for i in range(len(month_seq)):
            raw = vals[i] if i < len(vals) else None
            if raw is None or (isinstance(raw, float) and pd.isna(raw)):
                out.append(0.0)
            else:
                out.append(_round_headcount(float(raw)))
        return out

    data: dict[str, list[Any]] = {
        label_col: row_labels,
    }
    for i, key in enumerate(col_keys):
        data[key] = [
            phase_cells[i],
            month_labels[i],
            *(_eng_row(name)[i] for name in HC_ENGINEER_TYPES),
        ]
    data[comments_col] = [_comment_for(lbl) for lbl in row_labels]

    return pd.DataFrame(data)[[label_col, *col_keys, comments_col]]


def _format_hc_detail_task(task: dict[str, Any], *, index: int) -> str:
    """
    One computation line, e.g.
      1. "Regular Meeting / Discussion": 1x  0.1 HC / Manpower (per Month)
    """
    name = str(task.get("task") or f"Task {index}").strip() or f"Task {index}"
    n_hc = int(task.get("number_of_headcount") or 0)
    if "headcount_per_person" in task:
        per = float(task.get("headcount_per_person") or 0)
    elif n_hc > 0:
        per = float(task.get("headcount") or 0) / n_hc
    else:
        per = float(task.get("headcount") or 0)
    per_disp = f"{_round_headcount(per):.1f}"
    dist = str(task.get("distribution") or HC_DIST_PER_MONTH).strip()
    if dist == HC_DIST_PER_PHASE:
        dist_label = "per Phase"
    else:
        dist_label = "per Month"
    return (
        f'{index}. "{name}": {n_hc}x  {per_disp} HC / Manpower ({dist_label})'
    )


def build_headcount_detail_table(
    phases: list[str],
    tasks_by_phase: dict[str, dict[str, list[dict[str, Any]]]] | None = None,
) -> pd.DataFrame:
    """
    Detail how headcount is computed — one row per phase; tasks stacked
    as separate lines in each engineer cell:

      Phase | Structure Engineer | CAE Engineer | S&V Engineer
      Concept | 1. "Task A": …\\n2. "Task B": … | … | …
      BU | 1. "…" | … | …
    """
    ordered = [p for p in PHASES if p in {str(x) for x in phases}]
    if not ordered:
        ordered = [str(p) for p in phases]

    tasks_map = tasks_by_phase or {}
    rows: list[dict[str, Any]] = []
    for phase in ordered:
        eng_tasks = tasks_map.get(phase) or {}
        row: dict[str, Any] = {"Phase": phase}
        for eng in HC_ENGINEER_TYPES:
            tasks = list(eng_tasks.get(eng) or [])
            if not tasks:
                row[eng] = ""
            else:
                row[eng] = "\n".join(
                    _format_hc_detail_task(t, index=i + 1)
                    for i, t in enumerate(tasks)
                )
        rows.append(row)

    cols = ["Phase", *HC_ENGINEER_TYPES]
    return pd.DataFrame(rows, columns=cols) if rows else pd.DataFrame(columns=cols)


def export_headcount_xlsx(
    hc_df: pd.DataFrame,
    detail_df: pd.DataFrame | None = None,
) -> bytes:
    """
    Export headcount grid + computation detail as one workbook:
      Sheet "Headcount" — monthly HC table
      Sheet "Detail" — per-phase task formulas (newlines preserved)
    Light grey fill on Headcount Phase/Month rows and Detail header row.
    """
    fill_grey = PatternFill(
        start_color="D9D9D9", end_color="D9D9D9", fill_type="solid"
    )
    buf = BytesIO()
    hc = hc_df if hc_df is not None else pd.DataFrame()
    detail = detail_df if detail_df is not None else pd.DataFrame()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        hc.to_excel(writer, sheet_name="Headcount", index=False)
        detail.to_excel(writer, sheet_name="Detail", index=False)
        wb = writer.book
        # Headcount: modest column widths + grey Phase / Month rows
        ws_hc = wb["Headcount"]
        for col_idx in range(1, ws_hc.max_column + 1):
            letter = get_column_letter(col_idx)
            ws_hc.column_dimensions[letter].width = 14 if col_idx > 1 else 22
        for r in range(1, ws_hc.max_row + 1):
            label = str(ws_hc.cell(r, 1).value or "").strip()
            if label in ("Phase", "Month"):
                for c in range(1, ws_hc.max_column + 1):
                    ws_hc.cell(r, c).fill = fill_grey
        # Detail: wrap multiline task cells; grey header row
        ws_d = wb["Detail"]
        wrap = Alignment(wrap_text=True, vertical="top")
        for row in ws_d.iter_rows(
            min_row=1,
            max_row=ws_d.max_row,
            max_col=ws_d.max_column,
        ):
            for cell in row:
                cell.alignment = wrap
        if ws_d.max_row >= 1:
            for c in range(1, ws_d.max_column + 1):
                ws_d.cell(1, c).fill = fill_grey
        ws_d.column_dimensions["A"].width = 12
        for col_idx in range(2, ws_d.max_column + 1):
            ws_d.column_dimensions[get_column_letter(col_idx)].width = 48
        for r in range(2, ws_d.max_row + 1):
            # Grow row height for multiline cells
            max_lines = 1
            for col_idx in range(1, ws_d.max_column + 1):
                val = ws_d.cell(r, col_idx).value
                if val is None:
                    continue
                max_lines = max(max_lines, str(val).count("\n") + 1)
            ws_d.row_dimensions[r].height = max(18.0, 15.0 * max_lines + 6.0)
    return buf.getvalue()
